import io
import os
from typing import List, Tuple

import fitz  # PyMuPDF
import pytesseract
import streamlit as st
from docx import Document
from groq import Groq, APIError
from PIL import Image


APP_TITLE = "AI Study Buddy"
MODEL_NAME = "openai/gpt-oss-120b"
MAX_CONTEXT_CHARS = 12_000

st.set_page_config(
    page_title=APP_TITLE,
    page_icon="💜",
    layout="wide",
    initial_sidebar_state="expanded",
)

# Light-purple theme. Streamlit's built-in theme can be configured later too.
st.markdown(
    """
    <style>
      .stApp { background: #F8F5FF; color: #2D2142; }
      [data-testid="stSidebar"] { background: #EEE6FF; }
      h1, h2, h3 { color: #5B35A6; }
      div.stButton > button, div[data-testid="stFormSubmitButton"] > button {
        background: #7857D9; color: white; border: 0; border-radius: 12px;
      }
      div.stButton > button:hover, div[data-testid="stFormSubmitButton"] > button:hover {
        background: #603DBB; color: white; border: 0;
      }
      [data-testid="stChatMessage"] { border-radius: 14px; }
      .study-tag { color: #6941A5; font-weight: 600; }
    </style>
    """,
    unsafe_allow_html=True,
)


def get_api_key() -> str:
    """Read the key from an environment variable or Streamlit Secrets."""
    key = os.getenv("GROQ_API_KEY", "").strip()
    if key:
        return key
    try:
        return str(st.secrets.get("GROQ_API_KEY", "")).strip()
    except Exception:
        # st.secrets may not be configured during local/Colab development.
        return ""


def extract_pdf(file_bytes: bytes) -> str:
    parts = []
    with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
        for page_number, page in enumerate(pdf, start=1):
            page_text = page.get_text("text").strip()
            if page_text:
                parts.append(f"[PDF page {page_number}]\n{page_text}")
    return "\n\n".join(parts)


def extract_docx(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append(" | ".join(cell.text.strip() for cell in row.cells))
        if rows:
            parts.append(f"[Word table {table_index}]\n" + "\n".join(rows))
    return "\n\n".join(parts)


def extract_upload(uploaded_file) -> Tuple[str, str]:
    """Return extracted text and a short status message for one uploaded file."""
    name = uploaded_file.name
    suffix = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    file_bytes = uploaded_file.getvalue()

    try:
        if suffix == "pdf":
            text = extract_pdf(file_bytes)
            if not text.strip():
                return "", f"{name}: no selectable text found. If it is a scanned PDF, upload page photos for OCR for now."
            return text, f"{name}: PDF text extracted."
        if suffix == "docx":
            text = extract_docx(file_bytes)
            return text, f"{name}: Word text extracted."
        if suffix == "txt":
            text = file_bytes.decode("utf-8", errors="replace")
            return text, f"{name}: text file read."
        if suffix in {"png", "jpg", "jpeg"}:
            image = Image.open(io.BytesIO(file_bytes)).convert("RGB")
            text = pytesseract.image_to_string(image).strip()
            if not text:
                return "", f"{name}: OCR found no readable text. Try a clearer, well-lit image."
            return text, f"{name}: image OCR completed. Check mathematical symbols for accuracy."
        return "", f"{name}: unsupported file type."
    except Exception as exc:
        return "", f"{name}: could not read this file ({type(exc).__name__})."


def make_system_prompt(subject: str, level: str, mode: str, study_context: str) -> str:
    mode_instructions = {
        "Tutor chat": "Answer the student's question clearly and helpfully.",
        "Step-by-step solution": (
            "Solve the problem in numbered steps. Identify the given information, explain the method or formula, "
            "show intermediate working, state the final answer clearly, and check it when possible. For coding, "
            "explain the approach and then provide readable code. Do not invent missing values."
        ),
        "Generate MCQs": (
            "Create 5 multiple-choice questions. Give four options (A-D) per question, then provide an answer key "
            "with a short explanation for each answer. If the student names a topic, focus on that topic."
        ),
        "Summary": (
            "Create a useful study summary with headings, concise bullet points, key definitions or formulas, "
            "and a short list of important exam points. Do not add irrelevant material."
        ),
        "Flashcards": (
            "Create 8 concise flashcards in Question / Answer format. Focus on key concepts, definitions, "
            "formulas, or important code ideas."
        ),
    }

    context_instruction = (
        "The student uploaded the following study material. Use it as the main source when relevant. "
        "If the answer is not supported by the material, say so briefly and clearly distinguish any general explanation "
        "you add. Never pretend that the material says something it does not.\n\n"
        f"STUDY MATERIAL (may be shortened):\n{study_context[:MAX_CONTEXT_CHARS]}"
        if study_context.strip()
        else "No study material has been uploaded. Answer from your general knowledge and be clear if uncertain."
    )

    return (
        "You are AI Study Buddy, a patient tutor for Computer Science, Mathematics, and Statistics. "
        "Use simple, supportive language appropriate for the selected learning level. Show your reasoning as "
        "clear educational steps without skipping important calculations. For mathematics and statistics, define "
        "symbols and show formulas and substitutions. For Computer Science, use correct terminology and readable code. "
        "Do not fabricate facts, textbook quotations, or results.\n\n"
        f"Selected subject: {subject}. Learning level: {level}. Study mode: {mode}.\n"
        f"Task instructions: {mode_instructions[mode]}\n\n{context_instruction}"
    )


def ask_groq(messages: List[dict]) -> str:
    api_key = get_api_key()
    if not api_key:
        raise RuntimeError(
            "Groq API key not found. In Colab, load GROQ_API_KEY from Colab Secrets before starting Streamlit. "
            "When deployed, add GROQ_API_KEY in the app's Streamlit Secrets settings."
        )

    client = Groq(api_key=api_key)
    response = client.chat.completions.create(
        model=MODEL_NAME,
        messages=messages,
        temperature=0.3,
        max_completion_tokens=1400,
    )
    return response.choices[0].message.content or "I couldn't generate an answer. Please try again."


# Sidebar settings
with st.sidebar:
    st.markdown("# 💜 AI Study Buddy")
    st.caption("Learn smarter, one concept at a time.")
    subject = st.selectbox("Choose a subject", ["Computer Science", "Mathematics", "Statistics"])
    level = st.selectbox("Learning level", ["Beginner", "Intermediate", "Advanced", "Mixed levels"])
    mode = st.selectbox(
        "Study mode",
        ["Tutor chat", "Step-by-step solution", "Generate MCQs", "Summary", "Flashcards"],
    )
    st.divider()
    st.markdown("**Upload study material**")
    uploaded_files = st.file_uploader(
        "PDF, Word, TXT, or page photos",
        type=["pdf", "docx", "txt", "png", "jpg", "jpeg"],
        accept_multiple_files=True,
        help="For best OCR results, upload clear, straight, well-lit photos.",
    )
    st.caption("Your API key should be stored in Colab Secrets or Streamlit Secrets, never in this file.")


# Read uploaded files and combine their text for the current session.
extracted_parts = []
file_statuses = []
if uploaded_files:
    for uploaded_file in uploaded_files:
        extracted_text, status = extract_upload(uploaded_file)
        file_statuses.append(status)
        if extracted_text.strip():
            extracted_parts.append(f"[File: {uploaded_file.name}]\n{extracted_text}")

study_context = "\n\n---\n\n".join(extracted_parts)

st.title("💜 AI Study Buddy")
st.markdown("Your personal study partner for **Computer Science, Mathematics, and Statistics**.")
st.info("Choose a subject and study mode in the sidebar, then ask a question or request a study activity.")

if file_statuses:
    with st.expander("Uploaded file status", expanded=True):
        for status in file_statuses:
            st.write("• " + status)
        if study_context:
            st.caption(f"Extracted text is ready ({len(study_context):,} characters before context shortening).")
            with st.expander("Preview extracted text"):
                st.text(study_context[:4000] + ("\n\n[Preview shortened]" if len(study_context) > 4000 else ""))

if "chat_history" not in st.session_state:
    st.session_state.chat_history = []

col1, col2 = st.columns([5, 1])
with col2:
    if st.button("Clear chat"):
        st.session_state.chat_history = []
        st.rerun()

for item in st.session_state.chat_history:
    with st.chat_message(item["role"]):
        st.markdown(item["content"])

user_prompt = st.chat_input("Ask a question or name a chapter/topic...")
if user_prompt:
    st.session_state.chat_history.append({"role": "user", "content": user_prompt})
    with st.chat_message("user"):
        st.markdown(user_prompt)

    system_prompt = make_system_prompt(subject, level, mode, study_context)
    recent_history = st.session_state.chat_history[-7:]
    messages = [{"role": "system", "content": system_prompt}]
    messages.extend({"role": item["role"], "content": item["content"]} for item in recent_history)

    with st.chat_message("assistant"):
        with st.spinner("Your study buddy is thinking..."):
            try:
                answer = ask_groq(messages)
                st.markdown(answer)
                st.session_state.chat_history.append({"role": "assistant", "content": answer})
            except RuntimeError as exc:
                answer = str(exc)
                st.error(answer)
            except APIError as exc:
                answer = "The AI service returned an error. Check your API key, model availability, and Groq usage limits."
                st.error(answer)
                st.caption(f"Technical detail: {type(exc).__name__}")
            except Exception as exc:
                answer = "Something went wrong. Please try again after checking the setup."
                st.error(answer)
                st.caption(f"Technical detail: {type(exc).__name__}")

st.caption("Prototype note: private email allowlisting will be added before public deployment.")
